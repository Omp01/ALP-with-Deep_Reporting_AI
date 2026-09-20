"""
Test doubles for the ingestion pipeline's external dependencies.

These stand in for services that cannot run inside a test (an LLM, YouTube, object
storage). They are used ONLY by tests; the product has no equivalent fallback path.

`ScriptedAI` behaves like a well-behaved extractive model: it reads the material out of the
prompt it is given and answers with objectives and questions whose source quotes are real
sentences from that material. Tests can override any part of its answer to simulate a bad
model (ungrounded quotes, malformed JSON, an outage).
"""

import json
import re
from typing import Callable, Dict, List, Optional

import httpx

from shared.schemas.ai_provider import AICompletionRequest, AICompletionResponse, AIProvider, AIProviderError


def material_from_prompt(prompt: str) -> str:
    blocks = re.findall(r"<<<CONTENT_START[^>]*>>>\n(.*?)\n<<<CONTENT_END", prompt, flags=re.DOTALL)
    return "\n".join(blocks)


def sentences(text: str, minimum: int = 60) -> List[str]:
    found = [s.strip() for s in re.split(r"(?<=[.!?])\s+", re.sub(r"\s+", " ", text)) if len(s.strip()) >= minimum]
    seen, unique = set(), []
    for s in found:
        if s.lower() not in seen:
            seen.add(s.lower())
            unique.append(s)
    return unique


def extractive_analysis(material: str, competencies: Optional[List[dict]] = None) -> dict:
    facts = sentences(material)
    return {
        "summary": " ".join(facts[:2])[:600] or "This material introduces a technical topic in some detail.",
        "level": "beginner",
        "objectives": [f"Explain the idea that {s[0].lower() + s[1:]}"[:250].rstrip(".") for s in facts[:4]] or [
            "Describe the main idea of the material", "Apply the main idea to a simple example"],
        "concepts": [{"name": " ".join(s.split()[:3]), "explanation": s[:200]} for s in facts[:3]],
        "competencies": competencies or [
            {"name": "SQL Joins", "description": "Combine rows from several tables.", "domain": "sql",
             "bloom_level": "apply", "difficulty": 0.5, "existing_code": None}],
    }


def extractive_questions(material: str, count: int = 4, competency: str = "SQL Joins") -> dict:
    facts = sentences(material)
    questions = []
    for i, fact in enumerate(facts[:count]):
        distractors = [f"The material says the opposite: {facts[(i + k) % len(facts)][:60]} is false" for k in (1, 2, 3)]
        questions.append({
            "question": f"According to the material, which statement is accurate about: {' '.join(fact.split()[:6])}?",
            "options": [fact[:200]] + distractors, "correct_index": 0,
            "explanation": "This is stated directly in the material.", "difficulty": 0.4 + 0.1 * (i % 3),
            "competency": competency, "source_quote": fact, "chunk_index": 0,
        })
    return {"questions": questions}


def grading_fields(prompt: str) -> Dict[str, str]:
    """What the grading prompt tells the model: the skill code, the expected answer and the learner's answer."""
    def between(start: str, end: str) -> str:
        found = re.search(re.escape(start) + r"(.*?)" + re.escape(end), prompt, flags=re.DOTALL)
        return found.group(1).strip() if found else ""

    answer = re.search(r"<<<ANSWER_START id=[^>]*>>>\n(.*?)\n<<<ANSWER_END", prompt, flags=re.DOTALL)
    return {
        "skill_id": between("Competency code:", "\n"),
        "expected": between("EXPECTED ANSWER:\n", "\n\nRUBRIC:"),
        "answer": answer.group(1) if answer else "",
    }


def overlap_grade(fields: Dict[str, str]) -> dict:
    """
    A stand-in grader for tests: the share of the expected answer's longer words found in the learner's answer.
    Confident, quotes the start of the answer, names the skill it was given.
    """
    words = {w for w in re.findall(r"[a-z0-9]{4,}", fields["expected"].lower())}
    given = set(re.findall(r"[a-z0-9]{4,}", fields["answer"].lower()))
    signal = round(len(words & given) / len(words), 2) if words else 0.0
    return {
        "skill_id": fields["skill_id"], "correctness_signal": signal, "confidence": 0.9,
        "error_type": None if signal >= 0.7 else "knowledge_gap",
        "evidence_quote": " ".join(fields["answer"].split()[:8]), "feedback": "Compared with the expected points.",
        "rubric_scores": {},
    }


def reporting_package(prompt: str) -> dict:
    """The evidence package the reporting prompt carries."""
    found = re.search(r"<<<EVIDENCE_START id=[^>]*>>>\n(.*?)\n<<<EVIDENCE_END", prompt, flags=re.DOTALL)
    return json.loads(found.group(1)) if found else {}


def faithful_report(package: dict) -> dict:
    """A well-behaved reporting model: one claim per finding that connects it to its evidence, using only numbers from the package."""
    claims = [{"claim": f"Observed: {p['statement']}", "claim_type": p["claim_type"], "evidence_ids": p["evidence_ids"][:3], "metric_ids": p["metric_ids"], "confidence": 0.8}
              for p in package.get("patterns", [])[:3]]
    first = package.get("patterns", [{}])[0]
    return {"summary": first.get("statement", "Nothing to report."), "claims": claims, "limits": ["Only the supplied evidence was considered."]}


def faithful_items(prompt: str = "") -> dict:
    """Three statements for each self-report construct, the second one negatively worded."""
    out = []
    for construct in ("self_efficacy", "motivation", "self_regulation", "learning_anxiety"):
        out += [
            {"dimension": construct, "text": f"I feel good about how I am doing with {construct.replace('_', ' ')} in this subject.", "reverse": False},
            {"dimension": construct, "text": f"I often struggle with {construct.replace('_', ' ')} when the subject gets difficult.", "reverse": True},
            {"dimension": construct, "text": f"Most days I can handle {construct.replace('_', ' ')} for this course without trouble.", "reverse": False},
        ]
    return {"items": out}


def faithful_note(facts: dict) -> dict:
    quiz = facts.get("quiz", {})
    return {"note": f"You answered {quiz.get('correct')} of {quiz.get('total')} questions correctly in this check-in. Review the lessons listed in your report to firm up the ones you missed, and keep your study sessions short and regular."}


class ScriptedAI(AIProvider):
    """
    analysis / questions / grading: a dict, or a callable(material | fields) -> dict, or a raw string to return verbatim.
    A grading callable receives {"skill_id", "expected", "answer"} read from the prompt.
    error: an exception to raise from every call (simulates an outage).
    """

    def __init__(self, analysis=None, questions=None, error: Optional[Exception] = None, grading=None, reporting=None, items=None, note=None):
        self._analysis, self._questions, self._error, self._grading, self._reporting = analysis, questions, error, grading, reporting
        self._items, self._note = items, note
        self.calls: List[Dict[str, str]] = []

    async def complete(self, request: AICompletionRequest) -> AICompletionResponse:
        system, user = request.messages[0].content, request.messages[-1].content
        first_user = next(m.content for m in request.messages if m.role == "user")
        if "psychometric item writer" in system:
            self.calls.append({"kind": "items", "system": system, "user": first_user, "material": ""})
            if self._error is not None:
                raise self._error
            source = self._items if self._items is not None else faithful_items
            payload = source(first_user) if callable(source) else source
            text = payload if isinstance(payload, str) else json.dumps(payload)
            return AICompletionResponse(content=text, model="scripted-model", usage={}, finish_reason="stop", latency_ms=1)
        if "learning coach" in system:
            facts = json.loads(first_user.splitlines()[1])          # the prompt is: header line, the facts as one JSON line, footer line
            self.calls.append({"kind": "note", "system": system, "user": first_user, "material": "", "facts": facts})
            if self._error is not None:
                raise self._error
            source = self._note if self._note is not None else faithful_note
            payload = source(facts) if callable(source) else source
            text = payload if isinstance(payload, str) else json.dumps(payload)
            return AICompletionResponse(content=text, model="scripted-model", usage={}, finish_reason="stop", latency_ms=1)
        if "reporting analyst" in system:
            package = reporting_package(first_user)
            self.calls.append({"kind": "reporting", "system": system, "user": first_user, "material": "", "package": package})
            if self._error is not None:
                raise self._error
            source = self._reporting if self._reporting is not None else faithful_report
            payload = source(package) if callable(source) else source
            text = payload if isinstance(payload, str) else json.dumps(payload)
            return AICompletionResponse(content=text, model="scripted-model", usage={}, finish_reason="stop", latency_ms=1)
        if "You are a grading assistant" in system:
            fields = grading_fields(first_user)
            self.calls.append({"kind": "grading", "system": system, "user": first_user, "material": "", **fields})
            if self._error is not None:
                raise self._error
            source = self._grading if self._grading is not None else overlap_grade
            payload = source(fields) if callable(source) else source
            text = payload if isinstance(payload, str) else json.dumps(payload)
            return AICompletionResponse(content=text, model="scripted-model", usage={}, finish_reason="stop", latency_ms=1)
        material = material_from_prompt(first_user)
        kind = "analysis" if "instructional designer" in system else "questions"
        self.calls.append({"kind": kind, "system": system, "user": first_user, "material": material})
        if self._error is not None:
            raise self._error

        source = self._analysis if kind == "analysis" else self._questions
        if source is None:
            source = extractive_analysis if kind == "analysis" else extractive_questions
        payload = source(material) if callable(source) else source
        text = payload if isinstance(payload, str) else json.dumps(payload)
        return AICompletionResponse(content=text, model="scripted-model", usage={}, finish_reason="stop", latency_ms=1)

    async def health_check(self) -> bool:
        return True

    @property
    def provider_name(self) -> str:
        return "scripted-ai"


def outage() -> ScriptedAI:
    return ScriptedAI(error=AIProviderError("connection refused", provider="scripted-ai"))


class MemoryStorage:
    """In-memory replacement for the object-storage client."""

    def __init__(self):
        self.objects: Dict[str, bytes] = {}
        self.fail_uploads = False

    async def upload_file(self, object_key: str, data: bytes, content_type: str = "") -> str:
        if self.fail_uploads:
            raise RuntimeError("storage down")
        self.objects[object_key] = data
        return object_key

    async def download_file(self, object_key: str) -> bytes:
        if object_key not in self.objects:
            raise RuntimeError("NoSuchKey")
        return self.objects[object_key]

    async def delete_file(self, object_key: str) -> None:
        self.objects.pop(object_key, None)


# ---------------------------------------------------------------------------- YouTube
def watch_page(video_id: str, *, title="Joins explained", length="754", captions_url: Optional[str] = None,
               status="OK", reason=None, track_kind: Optional[str] = None) -> str:
    player = {
        "playabilityStatus": {"status": status, **({"reason": reason} if reason else {})},
        "videoDetails": {"title": title, "author": "Test Channel", "lengthSeconds": length, "shortDescription": "A lesson."},
    }
    if captions_url:
        track = {"baseUrl": captions_url, "languageCode": "en", "name": {"simpleText": "English"}}
        if track_kind:
            track["kind"] = track_kind
        player["captions"] = {"playerCaptionsTracklistRenderer": {"captionTracks": [track]}}
    return f"<html><script>var ytInitialPlayerResponse = {json.dumps(player)};var meta = 1;</script></html>"


def captions_json3(lines: List[str], gap_ms: int = 5000) -> dict:
    return {"events": [{"tStartMs": i * gap_ms, "segs": [{"utf8": line}]} for i, line in enumerate(lines)]}


def youtube_transport(video_id: str, *, oembed_status=200, captions: Optional[List[str]] = None, page_status="OK",
                      track_kind: Optional[str] = None, page_error=False, title="Joins explained") -> httpx.MockTransport:
    caption_url = f"https://www.youtube.com/api/timedtext?v={video_id}&lang=en" if captions is not None else None

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if "/oembed" in url:
            if oembed_status != 200:
                return httpx.Response(oembed_status, json={})
            return httpx.Response(200, json={"title": title, "author_name": "Test Channel", "thumbnail_url": "https://i.ytimg.com/x.jpg"})
        if "timedtext" in url:
            return httpx.Response(200, json=captions_json3(captions or []))
        if "/watch" in url:
            if page_error:
                return httpx.Response(500, text="boom")
            return httpx.Response(200, text=watch_page(video_id, captions_url=caption_url, status=page_status,
                                                       reason="Private video" if page_status != "OK" else None,
                                                       track_kind=track_kind, title=title))
        return httpx.Response(404)

    return httpx.MockTransport(handler)


# -------------------------------------------------------------------------- documents
PROSE = (
    "A relational database stores data in tables that are linked by keys. "
    "An INNER JOIN returns only the rows that have matching values in both tables being combined. "
    "A LEFT JOIN returns every row from the left table and the matching rows from the right table, "
    "and it fills the columns of unmatched rows with NULL values. "
    "Choosing the wrong kind of join is the most common reason a query silently drops or duplicates rows. "
    "The join condition states which columns must match, and it is usually written with the ON keyword. "
    "When a foreign key in one table refers to the primary key of another table, joining on those columns "
    "reconstructs the relationship between the records. "
    "A query planner may use an index on the joined columns to avoid scanning every row of both tables. "
    "Aggregating after a join, for example counting orders for each customer, requires a GROUP BY clause "
    "that names the customer columns being reported.\n\n"
    "Before writing a join, decide which table is the driving table and whether unmatched rows must be kept. "
    "If unmatched rows must be kept, use an outer join and remember that filters on the right table in a WHERE "
    "clause can accidentally turn the outer join back into an inner join. "
    "Testing a join on a tiny sample with known answers is the fastest way to catch a wrong assumption."
)


def build_pdf(paragraphs: List[str]) -> bytes:
    """A minimal but valid single-font PDF whose pages contain the given text (no external library needed)."""
    def wrap(text: str, width: int = 88) -> List[str]:
        words, lines, line = text.split(), [], ""
        for word in words:
            if len(line) + len(word) + 1 > width:
                lines.append(line)
                line = word
            else:
                line = f"{line} {word}".strip()
        return lines + ([line] if line else [])

    lines: List[str] = []
    for paragraph in paragraphs:
        lines += wrap(paragraph) + [""]
    per_page = 55
    pages = [lines[i:i + per_page] for i in range(0, len(lines), per_page)] or [[""]]

    objects: List[bytes] = []
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_obj = 3 + 2 * len(pages)
    for i, page_lines in enumerate(pages):
        content_obj = 4 + 2 * i
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_obj} 0 R "
            f"/Resources << /Font << /F1 {font_obj} 0 R >> >> >>".encode()
        )
        stream_lines = ["BT", "/F1 10 Tf", "12 TL", "40 760 Td"]
        for line in page_lines:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            stream_lines.append(f"({escaped}) Tj T*")
        stream_lines.append("ET")
        stream = "\n".join(stream_lines).encode("latin-1", "replace")
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def build_docx(paragraphs: List[str], heading: str = "Joins") -> bytes:
    import io
    import docx

    document = docx.Document()
    document.add_heading(heading, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def fake_mp4(size: int = 4096) -> bytes:
    return b"\x00\x00\x00\x18ftypmp42\x00\x00\x00\x00mp42isom" + b"\x00" * size


def fake_mp3(size: int = 2048) -> bytes:
    return b"ID3\x03\x00\x00\x00\x00\x00\x00" + b"\x00" * size
