"""
Document text extraction: real files in, real text out, and visible errors otherwise.
There is no fallback that turns unreadable bytes into "text".
"""

import io

import pytest

from shared.parsers import ExtractionError
from shared.parsers.text_extractor import TextExtractor
from tests.foundation.fakes import PROSE, build_docx, build_pdf

extract = TextExtractor.extract_from_bytes


def squash(text: str) -> str:
    return " ".join(text.split())


def test_plain_text_and_markdown_are_read_as_utf8():
    text, meta = extract("# Title\r\n\r\nCafé au lait.\r\n".encode("utf-8-sig"), "notes.md")
    assert text == "# Title\n\nCafé au lait." and meta["format"] == "text" and meta["character_count"] == len(text)


def test_text_that_is_not_utf8_is_an_error_not_mojibake():
    with pytest.raises(ExtractionError) as info:
        extract("Café".encode("latin-1"), "notes.txt")
    assert info.value.code == "invalid_encoding"


def test_a_pdf_yields_its_text_and_page_count():
    text, meta = extract(build_pdf(PROSE.split("\n\n")), "guide.pdf")
    assert "INNER JOIN returns only the rows that have matching values" in squash(text)
    assert meta["format"] == "pdf" and meta["page_count"] >= 1 and meta["character_count"] == len(text)


def test_a_multi_page_pdf_keeps_every_page():
    paragraphs = [f"Paragraph number {i} discusses topic {i} of the database course in reasonable detail." for i in range(90)]
    text, meta = extract(build_pdf(paragraphs), "long.pdf")
    assert meta["page_count"] >= 2
    assert "Paragraph number 0 " in text and "Paragraph number 89 " in text


def test_a_pdf_with_no_text_layer_returns_no_text():
    text, _ = extract(build_pdf([" "]), "scan.pdf")
    assert text == ""                       # the pipeline turns this into a clear "scanned document" failure


@pytest.mark.parametrize("data", [b"%PDF-1.4 this is not really a pdf", b"", b"\x00" * 200])
def test_a_broken_pdf_is_a_named_error(data):
    with pytest.raises(ExtractionError) as info:
        extract(data, "broken.pdf")
    assert info.value.code in ("corrupt_file", "missing_dependency")


def test_a_docx_keeps_headings_paragraphs_and_tables():
    import docx

    document = docx.Document()
    document.add_heading("Joining tables", level=1)
    document.add_heading("Inner joins", level=2)
    document.add_paragraph("An INNER JOIN returns matching rows.")
    table = document.add_table(rows=2, cols=2)
    for cell, value in zip(table._cells, ["Join", "Keeps", "INNER", "Matches only"]):
        cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)

    text, meta = extract(buffer.getvalue(), "guide.docx")
    assert text.startswith("# Joining tables\n\n## Inner joins\n\nAn INNER JOIN returns matching rows.")
    assert "Join | Keeps" in text and "INNER | Matches only" in text
    assert meta["format"] == "docx" and meta["paragraph_count"] >= 4


def test_a_docx_round_trip_of_the_sample_text():
    text, _ = extract(build_docx(PROSE.split("\n\n")), "guide.docx")
    assert squash(PROSE) in squash(text)


@pytest.mark.parametrize("data", [b"not a zip at all", b"PK\x03\x04garbage", b""])
def test_a_broken_docx_is_a_named_error(data):
    with pytest.raises(ExtractionError) as info:
        extract(data, "broken.docx")
    assert info.value.code == "corrupt_file"


def test_a_pptx_yields_slide_text_and_speaker_notes():
    pptx = pytest.importorskip("pptx")
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Join types"
    slide.placeholders[1].text = "INNER keeps matches only"
    slide.notes_slide.notes_text_frame.text = "Mention the NULL trap."
    buffer = io.BytesIO()
    presentation.save(buffer)

    text, meta = extract(buffer.getvalue(), "deck.pptx")
    assert "Slide 1" in text and "Join types" in text and "INNER keeps matches only" in text and "Notes: Mention the NULL trap." in text
    assert meta["slide_count"] == 1


def test_a_broken_pptx_is_a_named_error():
    with pytest.raises(ExtractionError) as info:
        extract(b"not a presentation", "deck.pptx")
    assert info.value.code in ("corrupt_file", "missing_dependency")


@pytest.mark.parametrize("name", ["movie.mp4", "sound.mp3", "old.doc", "program.exe", "noextension"])
def test_formats_that_are_not_documents_are_refused(name):
    with pytest.raises(ExtractionError) as info:
        extract(b"anything", name)
    assert info.value.code == "unsupported_format"
