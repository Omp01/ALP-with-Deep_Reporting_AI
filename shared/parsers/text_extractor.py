"""
Document text extraction.

Extracts real text from PDF, DOCX, PPTX and plain-text/Markdown files. It never invents
text: if a file cannot be read, or the library needed to read it is not installed,
`ExtractionError` is raised so the caller can report exactly what went wrong.

(An earlier version fell back to decoding the raw bytes as text on any error, which
turned a failed PPTX or PDF parse into "successful" garbage, and returned a hardcoded
transcript for audio and video. Both are gone.)

Audio and video are not handled here: they need speech-to-text, which lives in
`app.ingestion.transcription`.
"""

import io
import os
from typing import Any, Dict, Tuple

TEXT_EXTENSIONS = {"txt", "md", "markdown"}
DOCUMENT_EXTENSIONS = TEXT_EXTENSIONS | {"pdf", "docx", "pptx"}


class ExtractionError(Exception):
    """A file could not be turned into text. `code` is stable and machine-readable."""

    def __init__(self, message: str, code: str = "extraction_failed"):
        super().__init__(message)
        self.code = code


class TextExtractor:
    """Extracts raw text from supported document formats."""

    @staticmethod
    def extract_from_bytes(file_bytes: bytes, filename: str) -> Tuple[str, Dict[str, Any]]:
        ext = os.path.splitext(filename)[1].lower().lstrip(".")
        if ext in TEXT_EXTENSIONS:
            return TextExtractor._extract_text(file_bytes)
        if ext == "pdf":
            return TextExtractor._extract_pdf(file_bytes)
        if ext == "docx":
            return TextExtractor._extract_docx(file_bytes)
        if ext == "pptx":
            return TextExtractor._extract_pptx(file_bytes)
        raise ExtractionError(f"Cannot extract text from '.{ext}' files.", code="unsupported_format")

    @staticmethod
    def _extract_text(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
        try:
            text = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ExtractionError("The file is not valid UTF-8 text.", code="invalid_encoding") from exc
        text = text.replace("\r\n", "\n").strip()
        return text, {"format": "text", "character_count": len(text)}

    @staticmethod
    def _extract_pdf(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
        try:
            import fitz  # PyMuPDF (preferred: better layout handling)
        except ImportError:
            fitz = None

        pages = []
        try:
            if fitz is not None:
                doc = fitz.open(stream=file_bytes, filetype="pdf")
                if doc.needs_pass:
                    raise ExtractionError("The PDF is password protected.", code="encrypted")
                pages = [page.get_text() for page in doc]
                doc.close()
            else:
                try:
                    import pypdf
                except ImportError as exc:
                    raise ExtractionError(
                        "No PDF reader is installed (install PyMuPDF or pypdf).", code="missing_dependency"
                    ) from exc
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                if reader.is_encrypted:
                    raise ExtractionError("The PDF is password protected.", code="encrypted")
                pages = [page.extract_text() or "" for page in reader.pages]
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError(f"The PDF could not be read: {exc}", code="corrupt_file") from exc

        text = "\n\n".join(p.strip() for p in pages if p.strip()).strip()
        return text, {"format": "pdf", "page_count": len(pages), "character_count": len(text)}

    @staticmethod
    def _extract_docx(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
        try:
            import docx
        except ImportError as exc:
            raise ExtractionError("python-docx is not installed.", code="missing_dependency") from exc

        parts = []
        try:
            document = docx.Document(io.BytesIO(file_bytes))
            for paragraph in document.paragraphs:
                if paragraph.text.strip():
                    style = (paragraph.style.name or "") if paragraph.style is not None else ""
                    text = paragraph.text.strip()
                    # Keep heading structure so downstream chunking and analysis can see it.
                    if style.startswith("Heading"):
                        level = "".join(ch for ch in style if ch.isdigit()) or "1"
                        text = f"{'#' * min(int(level), 4)} {text}"
                    parts.append(text)
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        except Exception as exc:
            raise ExtractionError(f"The DOCX file could not be read: {exc}", code="corrupt_file") from exc

        text = "\n\n".join(parts).strip()
        return text, {"format": "docx", "paragraph_count": len(parts), "character_count": len(text)}

    @staticmethod
    def _extract_pptx(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise ExtractionError("python-pptx is not installed.", code="missing_dependency") from exc

        parts = []
        try:
            presentation = Presentation(io.BytesIO(file_bytes))
            slide_count = len(presentation.slides)
            for number, slide in enumerate(presentation.slides, 1):
                texts = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                        texts.append(shape.text_frame.text.strip())
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                    texts.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
                if texts:
                    parts.append(f"Slide {number}\n" + "\n".join(texts))
        except Exception as exc:
            raise ExtractionError(f"The PPTX file could not be read: {exc}", code="corrupt_file") from exc

        text = "\n\n".join(parts).strip()
        return text, {"format": "pptx", "slide_count": slide_count, "character_count": len(text)}
